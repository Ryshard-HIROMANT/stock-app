import os
import json
from flask import Flask, render_template, redirect, url_for, flash, request, send_file
from config import Config
from models import db, User, Material, Batch, Location, Transaction, Category, Transfer, TransferItem, TransferRequest, SpendingDraft, Revision
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from datetime import date, datetime
from io import BytesIO

login_manager = LoginManager()
login_manager.login_view = 'login'
login_manager.login_message = 'Пожалуйста, войдите для доступа.'


def create_app():
    app = Flask(__name__)
    app.config.from_object(Config)

    db.init_app(app)
    login_manager.init_app(app)

    with app.app_context():
        db.create_all()

    @login_manager.user_loader
    def load_user(user_id):
        return User.query.get(int(user_id))

    # ──────────────────────────────────────────
    # АВТОРИЗАЦИЯ
    # ──────────────────────────────────────────
    @app.route('/login', methods=['GET', 'POST'])
    def login():
        if current_user.is_authenticated:
            return redirect(url_for('index'))
        if request.method == 'POST':
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '')
            user = User.query.filter_by(username=username).first()
            if user and user.check_password(password):
                if user.is_active:
                    login_user(user)
                    flash(f'Добро пожаловать, {user.full_name or user.username}!', 'success')
                    next_page = request.args.get('next')
                    return redirect(next_page or url_for('index'))
                else:
                    flash('Ваша учётная запись отключена.', 'danger')
            else:
                flash('Неверное имя пользователя или пароль.', 'danger')
        return render_template('login.html')

    @app.route('/logout')
    @login_required
    def logout():
        logout_user()
        flash('Вы вышли из системы.', 'info')
        return redirect(url_for('login'))

    # ──────────────────────────────────────────
    # ГЛАВНАЯ СТРАНИЦА (дашборд)
    # ──────────────────────────────────────────
    @app.route('/')
    @login_required
    def index():
        today = date.today()
        batches = Batch.query.filter(
            Batch.is_active == True,
            Batch.quantity - Batch.used > 0
        ).order_by(Batch.expiry_date.asc()).all()
        total_materials = Material.query.count()
        total_batches = Batch.query.filter(Batch.is_active == True).count()
        expired_batches = Batch.query.filter(
            Batch.is_active == True, Batch.quantity - Batch.used > 0,
            Batch.expiry_date < today
        ).count()
        expiring_soon = Batch.query.filter(
            Batch.is_active == True, Batch.quantity - Batch.used > 0,
            Batch.expiry_date >= today,
            Batch.expiry_date <= db.func.date(today, '+90 days')
        ).count()
        reusable_count = Material.query.filter_by(is_reusable=True).count()
        pending_requests = TransferRequest.query.filter_by(status='pending').count()
        pending_drafts = SpendingDraft.query.filter_by(status='pending').count()
        return render_template('index.html',
                               batches=batches, total_materials=total_materials,
                               total_batches=total_batches, expired_batches=expired_batches,
                               expiring_soon=expiring_soon, reusable_count=reusable_count,
                               pending_requests=pending_requests, pending_drafts=pending_drafts,
                               today=today)

    # ──────────────────────────────────────────
    # ДЕТАЛИЗАЦИЯ
    # ──────────────────────────────────────────
    @app.route('/dashboard/expired')
    @login_required
    def dashboard_expired():
        today = date.today()
        batches = Batch.query.filter(
            Batch.is_active == True, Batch.quantity - Batch.used > 0,
            Batch.expiry_date < today
        ).order_by(Batch.expiry_date.asc()).all()
        return render_template('dashboard_filtered.html', batches=batches,
                               title='Просроченные позиции', filter_type='expired')

    @app.route('/dashboard/expiring')
    @login_required
    def dashboard_expiring():
        today = date.today()
        batches = Batch.query.filter(
            Batch.is_active == True, Batch.quantity - Batch.used > 0,
            Batch.expiry_date >= today,
            Batch.expiry_date <= db.func.date(today, '+90 days')
        ).order_by(Batch.expiry_date.asc()).all()
        return render_template('dashboard_filtered.html', batches=batches,
                               title='Истекают в течение 90 дней', filter_type='expiring')

    # ──────────────────────────────────────────
    # МАТЕРИАЛЫ
    # ──────────────────────────────────────────
    @app.route('/materials')
    @login_required
    def materials():
        search = request.args.get('search', '').strip()
        category_id = request.args.get('category_id', '').strip()
        reusable = request.args.get('reusable', '').strip()
        query = Material.query
        if search:
            query = query.filter(
                Material.name.ilike(f'%{search}%') |
                Material.size.ilike(f'%{search}%') |
                Material.barcode.ilike(f'%{search}%')
            )
        if category_id:
            query = query.filter(Material.category_id == int(category_id))
        if reusable == '1':
            query = query.filter(Material.is_reusable == True)
        elif reusable == '0':
            query = query.filter(Material.is_reusable == False)
        materials_list = query.order_by(Material.name).all()
        categories_list = Category.query.order_by(Category.name).all()
        return render_template('materials.html', materials=materials_list, search=search,
                               categories=categories_list, selected_category=category_id)

    @app.route('/materials/add', methods=['GET', 'POST'])
    @login_required
    def material_add():
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('materials'))
        if request.method == 'POST':
            name = request.form.get('name', '').strip()
            size = request.form.get('size', '').strip()
            barcode = request.form.get('barcode', '').strip()
            category_id = request.form.get('category_id', '').strip()
            is_reusable = request.form.get('is_reusable') == '1'
            min_stock = request.form.get('min_stock', '0').strip()
            note = request.form.get('note', '').strip()
            if not name:
                flash('Название обязательно!', 'danger')
                categories_list = Category.query.order_by(Category.name).all()
                return render_template('material_form.html', material=None, categories=categories_list)
            material = Material(
                name=name, size=size if size else None,
                barcode=barcode if barcode else None,
                category_id=int(category_id) if category_id else None,
                is_reusable=is_reusable,
                min_stock=int(min_stock) if min_stock.isdigit() else 0,
                note=note if note else None
            )
            db.session.add(material)
            db.session.commit()
            flash(f'Материал "{name}" добавлен!', 'success')
            return redirect(url_for('materials'))
        categories_list = Category.query.order_by(Category.name).all()
        return render_template('material_form.html', material=None, categories=categories_list)

    @app.route('/materials/<int:id>/edit', methods=['GET', 'POST'])
    @login_required
    def material_edit(id):
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('materials'))
        material = Material.query.get_or_404(id)
        if request.method == 'POST':
            material.name = request.form.get('name', '').strip()
            material.size = request.form.get('size', '').strip() or None
            material.barcode = request.form.get('barcode', '').strip() or None
            cat_id = request.form.get('category_id', '').strip()
            material.category_id = int(cat_id) if cat_id else None
            material.is_reusable = request.form.get('is_reusable') == '1'
            min_stock = request.form.get('min_stock', '0').strip()
            material.min_stock = int(min_stock) if min_stock.isdigit() else 0
            material.note = request.form.get('note', '').strip() or None
            db.session.commit()
            flash(f'Материал "{material.name}" обновлён!', 'success')
            return redirect(url_for('materials'))
        categories_list = Category.query.order_by(Category.name).all()
        return render_template('material_form.html', material=material, categories=categories_list)

    @app.route('/materials/<int:id>/delete', methods=['POST'])
    @login_required
    def material_delete(id):
        if not current_user.is_admin():
            flash('Только администратор.', 'danger')
            return redirect(url_for('materials'))
        material = Material.query.get_or_404(id)
        name = material.name
        for batch in material.batches:
            Transaction.query.filter_by(batch_id=batch.id).delete()
        Batch.query.filter_by(material_id=material.id).delete()
        db.session.delete(material)
        db.session.commit()
        flash(f'Материал "{name}" удалён.', 'info')
        return redirect(url_for('materials'))

    # ──────────────────────────────────────────
    # КАТЕГОРИИ
    # ──────────────────────────────────────────
    @app.route('/categories')
    @login_required
    def categories():
        cats = Category.query.order_by(Category.name).all()
        return render_template('categories.html', categories=cats)

    @app.route('/categories/add', methods=['GET', 'POST'])
    @login_required
    def category_add():
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('categories'))
        if request.method == 'POST':
            name = request.form.get('name', '').strip()
            description = request.form.get('description', '').strip()
            if not name:
                flash('Название обязательно!', 'danger')
                return render_template('category_form.html', category=None)
            if Category.query.filter_by(name=name).first():
                flash(f'Категория "{name}" уже существует!', 'danger')
                return render_template('category_form.html', category=None)
            cat = Category(name=name, description=description if description else None)
            db.session.add(cat)
            db.session.commit()
            flash(f'Категория "{name}" добавлена!', 'success')
            return redirect(url_for('categories'))
        return render_template('category_form.html', category=None)

    @app.route('/categories/<int:id>/edit', methods=['GET', 'POST'])
    @login_required
    def category_edit(id):
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('categories'))
        cat = Category.query.get_or_404(id)
        if request.method == 'POST':
            new_name = request.form.get('name', '').strip()
            existing = Category.query.filter(Category.name == new_name, Category.id != id).first()
            if existing:
                flash(f'Категория "{new_name}" уже существует!', 'danger')
                return render_template('category_form.html', category=cat)
            cat.name = new_name
            cat.description = request.form.get('description', '').strip() or None
            db.session.commit()
            flash(f'Категория "{cat.name}" обновлена!', 'success')
            return redirect(url_for('categories'))
        return render_template('category_form.html', category=cat)

    @app.route('/categories/<int:id>/delete', methods=['POST'])
    @login_required
    def category_delete(id):
        if not current_user.is_admin():
            flash('Только администратор.', 'danger')
            return redirect(url_for('categories'))
        cat = Category.query.get_or_404(id)
        Material.query.filter_by(category_id=cat.id).update({Material.category_id: None})
        db.session.delete(cat)
        db.session.commit()
        flash(f'Категория "{cat.name}" удалена.', 'info')
        return redirect(url_for('categories'))

    # ──────────────────────────────────────────
    # ЛОКАЦИИ
    # ──────────────────────────────────────────
    @app.route('/locations')
    @login_required
    def locations():
        locs = Location.query.order_by(Location.code).all()
        return render_template('locations.html', locations=locs)

    @app.route('/locations/add', methods=['GET', 'POST'])
    @login_required
    def location_add():
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('locations'))
        if request.method == 'POST':
            code = request.form.get('code', '').strip()
            description = request.form.get('description', '').strip()
            if not code:
                flash('Код локации обязателен!', 'danger')
                return render_template('location_form.html', location=None)
            if Location.query.filter_by(code=code).first():
                flash(f'Локация "{code}" уже существует!', 'danger')
                return render_template('location_form.html', location=None)
            loc = Location(code=code, description=description if description else None)
            db.session.add(loc)
            db.session.commit()
            flash(f'Локация "{code}" добавлена!', 'success')
            return redirect(url_for('locations'))
        return render_template('location_form.html', location=None)

    @app.route('/locations/<int:id>/edit', methods=['GET', 'POST'])
    @login_required
    def location_edit(id):
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('locations'))
        loc = Location.query.get_or_404(id)
        if request.method == 'POST':
            new_code = request.form.get('code', '').strip()
            existing = Location.query.filter(Location.code == new_code, Location.id != id).first()
            if existing:
                flash(f'Локация "{new_code}" уже существует!', 'danger')
                return render_template('location_form.html', location=loc)
            loc.code = new_code
            loc.description = request.form.get('description', '').strip() or None
            db.session.commit()
            flash(f'Локация "{loc.code}" обновлена!', 'success')
            return redirect(url_for('locations'))
        return render_template('location_form.html', location=loc)

    @app.route('/locations/<int:id>/delete', methods=['POST'])
    @login_required
    def location_delete(id):
        if not current_user.is_admin():
            flash('Только администратор.', 'danger')
            return redirect(url_for('locations'))
        loc = Location.query.get_or_404(id)
        Batch.query.filter_by(location_id=loc.id).update({Batch.location_id: None})
        db.session.delete(loc)
        db.session.commit()
        flash(f'Локация "{loc.code}" удалена.', 'info')
        return redirect(url_for('locations'))

    # ──────────────────────────────────────────
    # СКАНЕР
    # ──────────────────────────────────────────
    @app.route('/scanner')
    @login_required
    def scanner():
        return render_template('scanner.html')

    @app.route('/api/material-by-barcode')
    @login_required
    def api_material_by_barcode():
        barcode = request.args.get('barcode', '').strip()
        if not barcode:
            return {'found': False}
        material = Material.query.filter_by(barcode=barcode).first()
        if not material:
            return {'found': False}
        batches = Batch.query.filter(
            Batch.material_id == material.id, Batch.is_active == True,
            Batch.quantity - Batch.used > 0
        ).order_by(Batch.expiry_date.asc()).all()
        total_remaining = sum(b.remaining for b in batches)
        return {
            'found': True,
            'material': {
                'id': material.id, 'name': material.name, 'size': material.size,
                'barcode': material.barcode, 'category_id': material.category_id,
                'category_name': material.category_name, 'is_reusable': material.is_reusable,
            },
            'total_remaining': total_remaining,
        }

    @app.route('/api/categories')
    @login_required
    def api_categories():
        cats = Category.query.order_by(Category.name).all()
        return [{'id': c.id, 'name': c.name} for c in cats]

    # ──────────────────────────────────────────
    # ПРИХОД
    # ──────────────────────────────────────────
    @app.route('/operations/in', methods=['GET', 'POST'])
    @login_required
    def operation_in():
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        if request.method == 'POST':
            mat_id = request.form.get('material_id', '').strip()
            qty = int(request.form.get('quantity', 0))
            expiry_str = request.form.get('expiry_date', '').strip()
            location_code = request.form.get('location_code', '').strip()
            batch_number = request.form.get('batch_number', '').strip()
            note = request.form.get('note', '').strip()
            new_name = request.form.get('new_material_name', '').strip()
            new_size = request.form.get('new_material_size', '').strip()
            new_barcode = request.form.get('new_material_barcode', '').strip()
            new_category_id = request.form.get('new_material_category_id', '').strip()
            is_reusable = request.form.get('is_reusable') == '1'
            if qty <= 0:
                flash('Укажите количество больше 0.', 'danger')
                return redirect(url_for('operation_in'))
            if mat_id and mat_id != 'new':
                material = Material.query.get_or_404(int(mat_id))
            elif new_name:
                material = Material.query.filter_by(name=new_name, size=new_size if new_size else None).first()
                if not material:
                    material = Material(
                        name=new_name, size=new_size if new_size else None,
                        barcode=new_barcode if new_barcode else None,
                        category_id=int(new_category_id) if new_category_id else None,
                        is_reusable=is_reusable,
                    )
                    db.session.add(material)
                    db.session.flush()
            else:
                flash('Выберите материал или введите название нового.', 'danger')
                return redirect(url_for('operation_in'))
            location = Location.query.filter_by(code=location_code).first() if location_code else None
            expiry_date = datetime.strptime(expiry_str, '%Y-%m-%d').date() if expiry_str else None
            batch = Batch(
                material_id=material.id, location_id=location.id if location else None,
                batch_number=batch_number if batch_number else None,
                expiry_date=expiry_date, quantity=qty, used=0
            )
            db.session.add(batch)
            db.session.flush()
            t_type = 'reusable_in' if material.is_reusable else 'in'
            transaction = Transaction(
                user_id=current_user.id, batch_id=batch.id, type=t_type,
                quantity=qty, note=note if note else f'Приход "{material.name}"'
            )
            db.session.add(transaction)
            db.session.commit()
            flash(f'Приход: {material.name} — {qty} шт.!', 'success')
            return redirect(url_for('index'))
        materials_list = Material.query.order_by(Material.name).all()
        locations_list = Location.query.order_by(Location.code).all()
        categories_list = Category.query.order_by(Category.name).all()
        return render_template('operation_in.html',
                               materials=materials_list, locations=locations_list,
                               categories=categories_list, selected_material=None)

    # ──────────────────────────────────────────
    # РАСХОД (с корзиной)
    # ──────────────────────────────────────────
    @app.route('/operations/out', methods=['GET', 'POST'])
    @login_required
    def operation_out():
        if not current_user.can_create_draft():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))

        if request.method == 'POST':
            operation_id = request.form.get('operation_id', '').strip()
            doctor_name = request.form.get('doctor_name', '').strip()
            note = request.form.get('note', '').strip()
            cart_data = request.form.get('cart_data', '[]')
            cart = json.loads(cart_data)

            if not cart or not operation_id or not doctor_name:
                flash('Заполните все обязательные поля.', 'danger')
                return redirect(url_for('operation_out'))

            if current_user.role in ('admin', 'head', 'doctor_storekeeper', 'doctor'):
                for item in cart:
                    material = Material.query.get_or_404(int(item['id']))
                    qty = int(item['qty'])
                    batches = Batch.query.filter(
                        Batch.material_id == material.id, Batch.is_active == True,
                        Batch.quantity - Batch.used > 0
                    ).order_by(Batch.expiry_date.asc()).all()
                    total_available = sum(b.remaining for b in batches)
                    if qty > total_available:
                        flash(f'Недостаточно "{material.name}"! Доступно: {total_available} шт.', 'danger')
                        return redirect(url_for('operation_out'))
                    remaining_to_take = qty
                    for batch in batches:
                        if remaining_to_take <= 0:
                            break
                        take = min(batch.remaining, remaining_to_take)
                        batch.used += take
                        remaining_to_take -= take
                        t_type = 'reusable_out' if material.is_reusable else 'out'
                        db.session.add(Transaction(
                            user_id=current_user.id, batch_id=batch.id, type=t_type,
                            quantity=take, operation_id=operation_id, doctor_name=doctor_name,
                            note=note or f'Расход "{material.name}"'
                        ))
                db.session.commit()
                flash(f'Расход по операции {operation_id}: {len(cart)} позиций!', 'success')
            else:
                for item in cart:
                    draft = SpendingDraft(
                        user_id=current_user.id, material_id=int(item['id']),
                        quantity=int(item['qty']), operation_id=operation_id,
                        doctor_name=doctor_name, note=note
                    )
                    db.session.add(draft)
                db.session.commit()
                flash(f'Черновик по операции {operation_id}: {len(cart)} позиций. Ожидает врача.', 'info')

            return redirect(url_for('index'))

        materials_list = Material.query.order_by(Material.name).all()
        categories_list = Category.query.order_by(Category.name).all()
        return render_template('operation_out.html',
                               materials=materials_list, categories=categories_list)

    # ──────────────────────────────────────────
    # ПОДТВЕРЖДЕНИЕ СПИСАНИЙ
    # ──────────────────────────────────────────
    @app.route('/spending/confirm')
    @login_required
    def spending_confirm_list():
        if not current_user.can_confirm():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        drafts = SpendingDraft.query.filter_by(status='pending').order_by(SpendingDraft.created_at.desc()).all()
        return render_template('spending_confirm.html', drafts=drafts)

    @app.route('/spending/<int:id>/approve', methods=['POST'])
    @login_required
    def spending_approve(id):
        if not current_user.can_confirm():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        draft = SpendingDraft.query.get_or_404(id)
        material = draft.material
        batches = Batch.query.filter(
            Batch.material_id == material.id, Batch.is_active == True,
            Batch.quantity - Batch.used > 0
        ).order_by(Batch.expiry_date.asc()).all()
        total_available = sum(b.remaining for b in batches)
        if draft.quantity > total_available:
            flash(f'Недостаточно! Доступно: {total_available} шт.', 'danger')
            return redirect(url_for('spending_confirm_list'))
        remaining_to_take = draft.quantity
        for batch in batches:
            if remaining_to_take <= 0:
                break
            take = min(batch.remaining, remaining_to_take)
            batch.used += take
            remaining_to_take -= take
            db.session.add(Transaction(
                user_id=current_user.id, batch_id=batch.id, type='out',
                quantity=take, operation_id=draft.operation_id,
                doctor_name=draft.doctor_name,
                note=f'Подтверждено врачом: {draft.note or ""}'
            ))
        draft.status = 'confirmed'
        draft.confirmed_by = current_user.id
        draft.confirmed_at = datetime.utcnow()
        db.session.commit()
        flash(f'Списание подтверждено: {material.name} — {draft.quantity} шт.', 'success')
        return redirect(url_for('spending_confirm_list'))

    @app.route('/spending/<int:id>/reject', methods=['POST'])
    @login_required
    def spending_reject(id):
        if not current_user.can_confirm():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        draft = SpendingDraft.query.get_or_404(id)
        draft.status = 'rejected'
        draft.confirmed_by = current_user.id
        draft.confirmed_at = datetime.utcnow()
        db.session.commit()
        flash('Списание отклонено.', 'info')
        return redirect(url_for('spending_confirm_list'))

    # ──────────────────────────────────────────
    # ЗАЯВКИ НА ПЕРЕМЕЩЕНИЕ
    # ──────────────────────────────────────────
    @app.route('/requests')
    @login_required
    def requests_list():
        if current_user.can_create_request():
            reqs = TransferRequest.query.order_by(TransferRequest.created_at.desc()).all()
        else:
            reqs = TransferRequest.query.filter_by(from_user_id=current_user.id).order_by(TransferRequest.created_at.desc()).all()
        return render_template('requests.html', requests=reqs)

    @app.route('/requests/add', methods=['GET', 'POST'])
    @login_required
    def request_add():
        if not current_user.can_create_request():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        if request.method == 'POST':
            mat_id = request.form.get('material_id')
            qty = int(request.form.get('quantity', 0))
            note = request.form.get('note', '').strip()
            if not mat_id or qty <= 0:
                flash('Укажите материал и количество.', 'danger')
                return redirect(url_for('request_add'))
            req = TransferRequest(
                from_user_id=current_user.id, material_id=int(mat_id),
                quantity=qty, note=note, to_location_code='oper'
            )
            db.session.add(req)
            db.session.commit()
            flash('Заявка создана!', 'success')
            return redirect(url_for('requests_list'))
        materials_list = Material.query.order_by(Material.name).all()
        return render_template('request_form.html', materials=materials_list)

    # ──────────────────────────────────────────
    # ПЕРЕМЕЩЕНИЯ
    # ──────────────────────────────────────────
    @app.route('/transfers')
    @login_required
    def transfers_list():
        transfers = Transfer.query.order_by(Transfer.created_at.desc()).limit(50).all()
        return render_template('transfers.html', transfers=transfers)

    @app.route('/transfers/add', methods=['GET', 'POST'])
    @login_required
    def transfer_add():
        if not current_user.is_storekeeper():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        if request.method == 'POST':
            from_location_code = request.form.get('from_location', '').strip()
            to_location_code = request.form.get('to_location', '').strip()
            note = request.form.get('note', '').strip()
            if not from_location_code or not to_location_code:
                flash('Выберите обе локации.', 'danger')
                return redirect(url_for('transfer_add'))
            if from_location_code == to_location_code:
                flash('Локации должны быть разными.', 'danger')
                return redirect(url_for('transfer_add'))
            from_loc = Location.query.filter_by(code=from_location_code).first()
            to_loc = Location.query.filter_by(code=to_location_code).first()
            if not from_loc or not to_loc:
                flash('Локация не найдена.', 'danger')
                return redirect(url_for('transfer_add'))
            transfer = Transfer(
                from_location_id=from_loc.id, to_location_id=to_loc.id,
                user_id=current_user.id, note=note
            )
            db.session.add(transfer)
            db.session.flush()
            index = 0
            while True:
                batch_key = f'batch_id_{index}'
                qty_key = f'qty_{index}'
                if batch_key not in request.form:
                    break
                batch_id = request.form.get(batch_key)
                qty = int(request.form.get(qty_key, 0))
                if batch_id and qty > 0:
                    batch = Batch.query.get(int(batch_id))
                    if batch and batch.location_id == from_loc.id and batch.remaining >= qty:
                        batch.used += qty
                        new_batch = Batch(
                            material_id=batch.material_id, location_id=to_loc.id,
                            batch_number=batch.batch_number, expiry_date=batch.expiry_date,
                            quantity=qty, used=0
                        )
                        db.session.add(new_batch)
                        db.session.flush()
                        item = TransferItem(transfer_id=transfer.id, batch_id=batch.id, quantity=qty)
                        db.session.add(item)
                        db.session.add(Transaction(
                            user_id=current_user.id, batch_id=batch.id, type='move_out',
                            quantity=qty, note=f'Перемещение в {to_loc.code}'
                        ))
                        db.session.add(Transaction(
                            user_id=current_user.id, batch_id=new_batch.id, type='move_in',
                            quantity=qty, note=f'Перемещение из {from_loc.code}'
                        ))
                index += 1
            db.session.commit()
            flash('Перемещение выполнено!', 'success')
            return redirect(url_for('transfers_list'))
        locations_list = Location.query.order_by(Location.code).all()
        return render_template('transfer_form.html', locations=locations_list)

    @app.route('/api/batches-by-location')
    @login_required
    def api_batches_by_location():
        location_code = request.args.get('location', '').strip()
        if not location_code:
            return []
        location = Location.query.filter_by(code=location_code).first()
        if not location:
            return []
        batches = Batch.query.filter(
            Batch.location_id == location.id, Batch.is_active == True,
            Batch.quantity - Batch.used > 0
        ).order_by(Batch.expiry_date.asc()).all()
        return [{
            'id': b.id, 'name': b.material.name, 'size': b.material.size or '',
            'remaining': b.remaining,
            'expiry': b.expiry_date.strftime('%d.%m.%Y') if b.expiry_date else '—',
        } for b in batches]

    # ──────────────────────────────────────────
    # РЕВИЗИЯ
    # ──────────────────────────────────────────
    @app.route('/revision', methods=['GET', 'POST'])
    @login_required
    def revision():
        if not current_user.can_revision():
            flash('Недостаточно прав.', 'danger')
            return redirect(url_for('index'))
        if request.method == 'POST':
            batch_id = request.form.get('batch_id')
            new_quantity = int(request.form.get('new_quantity', 0))
            new_used = int(request.form.get('new_used', 0))
            note = request.form.get('note', '').strip()
            batch = Batch.query.get_or_404(int(batch_id))
            old_quantity = batch.quantity
            old_used = batch.used
            batch.quantity = new_quantity
            batch.used = new_used
            revision_record = Revision(
                user_id=current_user.id, batch_id=batch.id,
                old_quantity=old_quantity, new_quantity=new_quantity,
                old_used=old_used, new_used=new_used,
                note=f'Ревизия от {date.today().strftime("%d.%m.%Y")}. {note}'
            )
            db.session.add(revision_record)
            db.session.add(Transaction(
                user_id=current_user.id, batch_id=batch.id, type='revision',
                quantity=abs(new_quantity - old_quantity),
                note=f'Ревизия: qty {old_quantity}→{new_quantity}, used {old_used}→{new_used}'
            ))
            db.session.commit()
            flash('Ревизия проведена!', 'success')
            return redirect(url_for('revision'))
        location_code = request.args.get('location', '').strip()
        location = Location.query.filter_by(code=location_code).first() if location_code else None
        batches = []
        if location:
            batches = Batch.query.filter(
                Batch.location_id == location.id, Batch.is_active == True
            ).order_by(Batch.expiry_date.asc()).all()
        locations_list = Location.query.order_by(Location.code).all()
        return render_template('revision.html', locations=locations_list, batches=batches, selected_location=location_code)

    # ──────────────────────────────────────────
    # ЖУРНАЛ ОПЕРАЦИЙ
    # ──────────────────────────────────────────
    @app.route('/transactions')
    @login_required
    def transactions():
        transactions_list = Transaction.query.order_by(Transaction.created_at.desc()).limit(100).all()
        return render_template('transactions.html', transactions=transactions_list)

    # ──────────────────────────────────────────
    # СОЗДАНИЕ ТЕСТОВЫХ ПОЛЬЗОВАТЕЛЕЙ
    # ──────────────────────────────────────────
    @app.route('/init-users')
    def init_users():
        users_data = [
            ('admin', 'admin123', 'Администратор', 'admin'),
            ('zaved', 'zaved123', 'Заведующий отделением', 'head'),
            ('sestra', 'sestra123', 'Старшая медсестра', 'head_nurse'),
            ('vrach_sklad', 'vrach123', 'Врач-кладовщик', 'doctor_storekeeper'),
            ('vrach', 'vrach123', 'Врач-оператор', 'doctor'),
            ('xray', 'xray123', 'Рентгенлаборант', 'xray_lab'),
        ]
        for username, password, full_name, role in users_data:
            if not User.query.filter_by(username=username).first():
                user = User(username=username, full_name=full_name, role=role)
                user.set_password(password)
                db.session.add(user)
        db.session.commit()
        return 'Пользователи созданы! <a href="/">Войти</a>'

    # ──────────────────────────────────────────
    # ЭКСПОРТ В EXCEL
    # ──────────────────────────────────────────
    @app.route('/export/excel')
    @login_required
    def export_excel():
        if not current_user.can_reports():
            flash('Недостаточно прав для экспорта.', 'danger')
            return redirect(url_for('index'))

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = Workbook()
        ws_default = wb.active
        ws_default.title = "Пусто"
        first_sheet = True

        today = date.today()
        locations = Location.query.order_by(Location.code).all()

        header_fill = PatternFill(start_color='0d6efd', end_color='0d6efd', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True, size=11)
        cat_fill = PatternFill(start_color='cfe2ff', end_color='cfe2ff', fill_type='solid')
        cat_font = Font(bold=True, size=11)
        red_fill = PatternFill(start_color='ffe0e0', end_color='ffe0e0', fill_type='solid')
        yellow_fill = PatternFill(start_color='fff3cd', end_color='fff3cd', fill_type='solid')
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        headers = ['Наименование', 'Размер', 'Штрихкод', 'Партия',
                   'Срок годности', 'Приход', 'Расход', 'Остаток', 'Статус']

        for loc in locations:
            batches = Batch.query.filter(
                Batch.location_id == loc.id,
                Batch.is_active == True
            ).order_by(Batch.expiry_date.asc()).all()

            if not batches:
                continue

            sheet_name = f'{loc.code} - {loc.description}'[:31]
            if first_sheet:
                ws = ws_default
                ws.title = sheet_name
                first_sheet = False
            else:
                ws = wb.create_sheet(title=sheet_name)

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin_border

            categories_dict = {}
            for b in batches:
                cat_name = b.material.category_name or 'Без категории'
                if cat_name not in categories_dict:
                    categories_dict[cat_name] = []
                categories_dict[cat_name].append(b)

            row = 2
            for cat_name in sorted(categories_dict.keys()):
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
                cell = ws.cell(row=row, column=1, value=f'{cat_name}')
                cell.fill = cat_fill
                cell.font = cat_font
                cell.border = thin_border
                row += 1

                for b in categories_dict[cat_name]:
                    remaining = b.remaining
                    if remaining == 0 and b.quantity == 0:
                        continue
                    status = 'Активен'
                    fill = None
                    if b.expiry_date and b.expiry_date < today:
                        status = 'Просрочен'
                        fill = red_fill
                    elif b.expiry_date and (b.expiry_date - today).days <= 90:
                        status = 'Истекает'
                        fill = yellow_fill

                    data = [
                        b.material.name,
                        b.material.size or '',
                        b.material.barcode or '',
                        b.batch_number or '',
                        b.expiry_date.strftime('%d.%m.%Y') if b.expiry_date else '',
                        b.quantity,
                        b.used,
                        remaining,
                        status,
                    ]
                    for col, value in enumerate(data, 1):
                        cell = ws.cell(row=row, column=col, value=value)
                        cell.border = thin_border
                        if fill:
                            cell.fill = fill
                    row += 1

            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = min(max_length + 2, 40)

        # Удаляем пустой лист, если есть другие
        if len(wb.sheetnames) > 1:
            wb.remove(wb["Пусто"])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'Остатки_по_складам_{today.strftime("%d.%m.%Y")}.xlsx'
        return send_file(output, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True)
    # ──────────────────────────────────────────
    # УПРАВЛЕНИЕ ПОЛЬЗОВАТЕЛЯМИ (админ)
    # ──────────────────────────────────────────
    @app.route('/admin/users')
    @login_required
    def admin_users():
        if not current_user.is_admin():
            flash('Только для администратора.', 'danger')
            return redirect(url_for('index'))
        users = User.query.order_by(User.role, User.username).all()
        return render_template('admin_users.html', users=users)

    @app.route('/admin/users/add', methods=['GET', 'POST'])
    @login_required
    def admin_user_add():
        if not current_user.is_admin():
            flash('Только для администратора.', 'danger')
            return redirect(url_for('index'))
        if request.method == 'POST':
            username = request.form.get('username', '').strip()
            password = request.form.get('password', '').strip()
            full_name = request.form.get('full_name', '').strip()
            role = request.form.get('role', '').strip()
            if not username or not password or not role:
                flash('Логин, пароль и роль обязательны!', 'danger')
                return render_template('admin_user_form.html', user=None)
            if User.query.filter_by(username=username).first():
                flash(f'Логин "{username}" уже занят!', 'danger')
                return render_template('admin_user_form.html', user=None)
            user = User(username=username, full_name=full_name, role=role)
            user.set_password(password)
            db.session.add(user)
            db.session.commit()
            flash(f'Пользователь "{username}" создан!', 'success')
            return redirect(url_for('admin_users'))
        return render_template('admin_user_form.html', user=None)

    @app.route('/admin/users/<int:id>/edit', methods=['GET', 'POST'])
    @login_required
    def admin_user_edit(id):
        if not current_user.is_admin():
            flash('Только для администратора.', 'danger')
            return redirect(url_for('index'))
        user = User.query.get_or_404(id)
        if request.method == 'POST':
            new_username = request.form.get('username', '').strip()
            existing = User.query.filter(User.username == new_username, User.id != id).first()
            if existing:
                flash(f'Логин "{new_username}" уже занят!', 'danger')
                return render_template('admin_user_form.html', user=user)
            user.username = new_username
            user.full_name = request.form.get('full_name', '').strip()
            user.role = request.form.get('role', '').strip()
            password = request.form.get('password', '').strip()
            if password:
                user.set_password(password)
            user.is_active = request.form.get('is_active') == '1'
            db.session.commit()
            flash(f'Пользователь "{user.username}" обновлён!', 'success')
            return redirect(url_for('admin_users'))
        return render_template('admin_user_form.html', user=user)

    @app.route('/admin/users/<int:id>/delete', methods=['POST'])
    @login_required
    def admin_user_delete(id):
        if not current_user.is_admin():
            flash('Только для администратора.', 'danger')
            return redirect(url_for('index'))
        if id == current_user.id:
            flash('Нельзя удалить самого себя!', 'danger')
            return redirect(url_for('admin_users'))
        user = User.query.get_or_404(id)
        username = user.username
        db.session.delete(user)
        db.session.commit()
        flash(f'Пользователь "{username}" удалён.', 'info')
        return redirect(url_for('admin_users'))
    # ──────────────────────────────────────────
    # API: список врачей
    # ──────────────────────────────────────────
    @app.route('/api/doctors')
    @login_required
    def api_doctors():
        doctors = User.query.filter(
            User.role.in_(['doctor', 'doctor_storekeeper', 'head']),
            User.is_active == True
        ).order_by(User.full_name).all()
        return [{'id': u.id, 'name': u.full_name or u.username} for u in doctors]
    
    return app


if __name__ == '__main__':
    app = create_app()
    port = int(os.environ.get('PORT', 5050))
    app.run(debug=False, host='0.0.0.0', port=port)